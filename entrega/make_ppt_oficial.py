"""Arma la presentación sobre la plantilla oficial del Qiskit Fall Fest 2026.

Uso:
    python entrega/make_ppt_oficial.py "<plantilla.pptx>" [entrega/QGate_oficial.pptx]

Respeta la primera y la última slide de la plantilla (solo llena título, nombres y equipo) y reemplaza
las slides intermedias por las slides de contenido de QGate (entrega/slides/slide02..08.png) con sus
notas del orador tomadas de entrega/QGate.pptx. Correr antes: python entrega/make_ppt.py
"""
from __future__ import annotations

import copy
import json
import os
import sys

from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches

HERE = os.path.dirname(os.path.abspath(__file__))
TITLE = ("QGate: Flight-to-Gate Assignment with", "Hybrid QAOA for the Qiskit Fall Fest 2026")


def set_text(shape, text):
    """Reemplaza el texto conservando el formato del primer run."""
    tf = shape.text_frame
    runs = [r for p in tf.paragraphs for r in p.runs]
    if not runs:
        return
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""


def walk(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == 6:
            yield from walk(sh.shapes)


def main():
    tpl = sys.argv[1]
    out = sys.argv[2] if len(sys.argv) > 2 else os.path.join(HERE, "QGate_oficial.pptx")
    team = json.load(open(os.path.join(HERE, "equipo.json"), encoding="utf-8"))
    ours = Presentation(os.path.join(HERE, "QGate.pptx"))
    notes = [s.notes_slide.notes_text_frame.text if s.has_notes_slide else "" for s in ours.slides]
    prs = Presentation(tpl)

    # slide 1: título en inglés, nombres; descripciones vacías (no inventamos roles)
    first = prs.slides[0]
    names = iter(team["integrantes"])
    for sh in walk(first.shapes):
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text
        if "TITULO" in t:
            paras = [p for p in sh.text_frame.paragraphs if p.runs]
            paras[0].runs[0].text = TITLE[0]
            for r in paras[0].runs[1:]:
                r.text = ""
            paras[1].runs[0].text = TITLE[1]
            for r in paras[1].runs[1:]:
                r.text = ""
        elif "<Nombre>" in t:
            set_text(sh, next(names, ""))
            grow = Inches(3.2)                      # la caja de la plantilla es para "<Nombre>"
            sh.left, sh.width = sh.left - grow, sh.width + grow
            for para in sh.text_frame.paragraphs:
                para.alignment = PP_ALIGN.RIGHT
        elif "<Descripcion breve>" in t:
            set_text(sh, "")
    # última slide: nombre del equipo
    last = prs.slides[len(prs.slides) - 1]
    for sh in walk(last.shapes):
        if sh.has_text_frame and "TEAM NOMBRE" in sh.text_frame.text:
            set_text(sh, f"TEAM {team['equipo'].upper()}")

    # quitar las slides intermedias de la plantilla
    sld = prs.slides._sldIdLst
    ids = list(sld)
    for el in ids[1:-1]:
        prs.part.drop_rel(el.rId)
        sld.remove(el)
    # renombrar la parte de la última slide para que no choque con las nuevas (slideN.xml)
    from pptx.opc.packuri import PackURI
    last.part.partname = PackURI("/ppt/slides/slide999.xml")
    # agregar las slides de contenido (imágenes 16:9 a pantalla completa) y mover la última al final
    blank = next(l for l in prs.slide_layouts if l.name == "Blank")
    for i in range(2, 9):
        png = os.path.join(HERE, "slides", f"slide{i:02d}.png")
        s = prs.slides.add_slide(blank)
        s.shapes.add_picture(png, 0, 0, width=prs.slide_width, height=prs.slide_height)
        if i - 1 < len(notes):
            s.notes_slide.notes_text_frame.text = notes[i - 1]
    last.notes_slide.notes_text_frame.text = ("Desarrollado con asistencia de Claude Code (Anthropic), supervisada por el equipo; parte del motor se preparó antes de la jornada; los cuatro integrantes pueden explicar su funcionamiento. Créditos: Qiskit, Qiskit Aer, qiskit-ibm-runtime, "
                                              "NumPy, SciPy, MIPLIB 2017 (qap10), Claude Code.")
    last_el = list(sld)[1]
    sld.remove(last_el)
    sld.append(last_el)
    prs.save(out)
    print(out, len(prs.slides), "slides")


if __name__ == "__main__":
    main()
